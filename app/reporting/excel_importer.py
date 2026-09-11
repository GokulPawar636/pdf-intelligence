"""Map explicitly labelled Excel measurements into the shared report layout."""
from pathlib import Path
from decimal import Decimal, InvalidOperation
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from .summary_exporter import SummaryRow, number, tolerance_limits, export_summary


ALIASES = {
    'object': {'object', 'feature', 'measured feature object', 'object name'},
    'control': {'control', 'characteristic', 'measurement type control', 'measurement type'},
    'nominal': {'nominal', 'nom', 'nominal value'},
    'tolerance': {'tolerance', 'tol'},
    'unit': {'unit', 'units'},
    'lower': {'lower limit', 'lsl', 'lower tolerance'},
    'upper': {'upper limit', 'usl', 'upper tolerance'},
}


def normalize(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').lower()).strip()


def read_excel(path):
    path = Path(path)
    if path.suffix.lower() not in ('.xlsx', '.xls'):
        raise ValueError('Upload an .xlsx or .xls workbook.')
    if path.stat().st_size > 100*1024*1024:
        raise ValueError('Excel input exceeds 100 MB.')
    if path.suffix.lower() == '.xlsx':
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 500*1024*1024:
                raise ValueError('Expanded workbook exceeds 500 MB.')
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if sum(s.max_row*s.max_column for s in workbook) > 2_000_000:
                raise ValueError('Workbook exceeds the 2 million cell processing limit.')
            sheets = [(s.title, list(s.values)) for s in workbook]
        finally:
            workbook.close()
    else:
        import xlrd
        workbook = xlrd.open_workbook(path)
        try:
            if sum(s.nrows*s.ncols for s in workbook.sheets()) > 2_000_000:
                raise ValueError('Workbook exceeds the 2 million cell processing limit.')
            sheets = [(s.name, [s.row_values(r) for r in range(s.nrows)]) for s in workbook.sheets()]
        finally:
            workbook.release_resources()
    return sheets


def numeric(value):
    if isinstance(value, bool) or value in (None, ''):
        return None
    try:
        result = Decimal(str(value).strip())
        if not result.is_finite() or abs(result) > Decimal('1e100') or len(result.as_tuple().digits) > 15:
            return None
        return result
    except InvalidOperation:
        return None


def validate_mapping(proposal, values):
    r, columns, readings = proposal.get('header_row'), proposal.get('columns'), proposal.get('readings')
    if type(r) is not int or not 0 <= r < len(values) or not isinstance(columns, dict) or set(columns) != set(ALIASES):
        raise ValueError('Invalid header row or field mapping.')
    if columns['object'] is None or columns['control'] is None or not isinstance(readings, list) or not readings:
        raise ValueError('Choose feature, control, and at least one reading column.')
    indices = [i for i in columns.values() if i is not None]+readings
    if any(type(i) is not int or not 0 <= i < len(values[r]) for i in indices) or len(set(indices)) != len(indices):
        raise ValueError('Mapped columns must exist and cannot overlap.')


def formula_warnings(path):
    if path.suffix.lower() != '.xlsx':
        return []
    count = 0
    ns = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name):
                continue
            for cell in ET.fromstring(archive.read(name)).findall('.//m:c', ns):
                if cell.find('m:f', ns) is None:
                    continue
                count += 1
                cached = cell.find('m:v', ns)
                if cached is None or cached.text is None:
                    raise ValueError(f'Source formula has no saved result: {name} {cell.get("r")}. Open the source in Excel, recalculate, save, and upload again.')
    return [f'{count} source formula(s): saved results were imported. Verify the source was recalculated before upload.'] if count else []


def import_excel(path, output, *, strict=False, use_ai=False, mappings=None):
    path, output = Path(path), Path(output)
    sheets = read_excel(path)
    rows, headers, warnings = [], [], formula_warnings(path)
    basic = {}
    for sheet_name, values in sheets:
        proposal = (mappings or {}).get(sheet_name)
        if proposal:
            validate_mapping(proposal, values)
        if not proposal and use_ai and not any(any(normalize(v) in ALIASES['object'] for v in r) and
                              any(normalize(v) in ALIASES['control'] for v in r) for r in values):
            from app.semantic.excel_mapping import suggest_mapping
            proposal = suggest_mapping(values)
            if proposal:
                warnings.append(f'{sheet_name}: Groq proposed header mappings. Verify mapped columns against the source.')
        mapping = readings = None
        header_offsets = {}
        local_rows = []
        for row_index, values_row in enumerate(values, 1):
            names = [normalize(v) for v in values_row]
            if mapping is None and len(values_row) > 1 and values_row[1] not in (None, ''):
                if names[0] in ('part number','part name','part no','customer','date','operator','table type','drawing number','revision'):
                    basic.setdefault(str(values_row[0]), str(values_row[1]))
            candidate = {key: next((i for i,n in enumerate(names) if n in aliases), None)
                         for key,aliases in ALIASES.items()}
            ai_header = proposal is not None and row_index == proposal['header_row']+1
            if ai_header:
                candidate = proposal['columns']
            if ai_header or (proposal is None and candidate['object'] is not None and candidate['control'] is not None):
                mapping = candidate
                # Only explicitly named readings or date/sample report headers qualify.
                readings = [i for i,v in enumerate(values_row) if
                            re.match(r'^(reading|measurement|meas|actual|sample)(?:\s|\d|$)', names[i])
                            and i not in mapping.values() or
                            re.match(r'^\d{1,4}[./-]\d{1,2}[./-]\d{2,4}', str(v or ''))]
                if ai_header:
                    readings = proposal['readings']
                if not readings:
                    raise ValueError(f'{sheet_name}, row {row_index}: no identifiable reading columns. Use Reading 1, Reading 2, or dated sample headers.')
                signature = (tuple(mapping.items()), tuple((i,str(values_row[i])) for i in readings))
                if signature in header_offsets:
                    offset = header_offsets[signature]
                    continue
                offset = len(headers)
                header_offsets[signature] = offset
                for i in readings:
                    original = str(values_row[i])
                    date = re.match(r'\d{1,4}[./-]\d{1,2}[./-]\d{2,4}', original)
                    sample = re.search(r'NO_?\s*(\w+)', original, re.I)
                    headers.append({'title': f'Measurement {len(headers)+1}',
                                    'label': (date[0] + ('\nSample '+sample[1] if sample else '')) if date else original[:80],
                                    'source': f'{path.name} / {sheet_name} / {original}'})
                continue
            if mapping is None or not any(v not in (None, '') for v in values_row):
                continue
            def value(key):
                i = mapping[key]
                return values_row[i] if i is not None and i < len(values_row) else None
            raw = [values_row[i] if i < len(values_row) else None for i in readings]
            if not any(v not in (None,'') for v in raw):
                warnings.append(f'{sheet_name} row {row_index}: no cached readings; row not included. Recalculate source formulas in Excel if applicable.')
                continue
            measured = [numeric(v) for v in raw]
            if any(v not in (None,'') and n is None for v,n in zip(raw,measured)):
                raise ValueError(f'{sheet_name} row {row_index}: nonnumeric reading; correct the source before export.')
            if not value('object') or not value('control'):
                raise ValueError(f'{sheet_name} row {row_index}: missing feature or control identity.')
            nominal = numeric(value('nominal'))
            tol = str(value('tolerance') or '')
            lower, upper = tolerance_limits(nominal, tol)
            explicit = value('lower') is not None or value('upper') is not None
            if explicit:
                lower, upper = numeric(value('lower')), numeric(value('upper'))
            notes = []
            for field in ('nominal','lower','upper'):
                if value(field) not in (None,'') and numeric(value(field)) is None:
                    raise ValueError(f'{sheet_name} row {row_index}: invalid {field} specification; correct the source.')
            if explicit:
                derived_lower, derived_upper = tolerance_limits(nominal, tol)
                if ((lower is not None and derived_lower is not None and lower != derived_lower) or
                    (upper is not None and derived_upper is not None and upper != derived_upper)):
                    notes.append('Explicit limits conflict with the tolerance; verify the source specification.')
            if lower is None or upper is None or lower > upper:
                notes.append('Specification limits require review. Missing nominal values or corrupted/unsigned tolerances are not inferred.')
            unit = str(value('unit') or '')
            if 'angle' in str(value('control')).lower() and unit.lower() in ('inches','inch','mm','millimeters'):
                notes.append('Source gives a length unit for an angle. Verify the unit.')
            if not unit:
                notes.append('Source unit is not stated; verify before comparison.')
            source = f'{path.name} / {sheet_name} row {row_index}'
            row = SummaryRow(str(value('object')), str(value('control')), nominal, lower, upper, tol,
                             [None]*offset+measured, page=f'{sheet_name}!{row_index}', unit=unit,
                             warnings=notes, evidence=source+'\n'+str(values_row), explicit_limits=explicit)
            balloon = re.search(r'Ballooning\s+No\.?\s*(\d+[A-Za-z]?(?:,\d+[A-Za-z]?)*)', row.object_name, re.I)
            row.balloon = balloon[1] if balloon else ''
            local_rows.append(row)
        if not local_rows and any(any(v not in (None,'') for v in r) for r in values):
            warnings.append(f'{sheet_name}: no recognizable measurement table; sheet not included.')
        rows.extend(local_rows)
    if not rows:
        raise ValueError('No recognizable measurement table. Required headers: Object/Feature, Control, and Reading/Sample or dated measurement columns.')
    for row in rows:
        row.measurements += [None]*(len(headers)-len(row.measurements))
    warnings += list(dict.fromkeys(note for row in rows for note in row.warnings))
    if strict and warnings:
        raise ValueError('Quality checks blocked export: '+' | '.join(warnings[:5]))
    export_summary(None, None, output, summary_rows=rows, measurement_headers=headers,
                   logo_path=Path(__file__).resolve().parents[1]/'assets'/'company_logo.jpg',
                   numeric_format='0.'+'0'*max(3, min(10, max(-v.as_tuple().exponent for row in rows for v in row.measurements if v is not None))),
                   report_title=f'Inspection Summary - {path.stem}',
                   source_caption=f'Source: {path.name}. Measurement columns follow source order. Blank readings remain blank.',
                   title_fields_override=list(basic.items()) + [('Characteristics',str(len(rows))),('Measurement columns',str(len(headers))),
                                          ('Source units', ', '.join(sorted({r.unit for r in rows if r.unit})) or 'Not stated'),
                                          ('Import status','Review required' if warnings else 'Checks passed')])
    return {'name': output.name, 'contents': output.read_bytes(), 'warnings': warnings,
            'measurement_columns': len(headers), 'files': 1, 'pages': len(sheets), 'source': path.name}
