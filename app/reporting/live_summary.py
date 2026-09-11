"""Excel formulas for editable measurement reports (no macros required)."""
from copy import copy
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill


def reading_span(row, mean_column, sheet_prefix=""):
    # Anchors lie OUTSIDE the input area, so inserting before the first reading
    # or immediately before Mean expands the range after Excel adjusts references.
    return (f"INDEX({sheet_prefix}${row}:${row},1,COLUMN({sheet_prefix}$F{row})+1):"
            f"INDEX({sheet_prefix}${row}:${row},1,COLUMN({sheet_prefix}${letter(mean_column)}{row})-1)")


def offset_formula(row, upper=False):
    text = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(F{row}," ",""),"\u2212","-"),"+/-","\u00b1")'
    left = f'LEFT({text},FIND("/",{text})-1)'
    right = f'MID({text},FIND("/",{text})+1,255)'
    signed = lambda part: f'OR(LEFT({part},1)="+",LEFT({part},1)="-",VALUE({part})=0)'
    pair = f'IF(AND({signed(left)},{signed(right)}),{"MAX" if upper else "MIN"}(VALUE({left}),VALUE({right})),"")'
    return (f'=IFERROR(IF(LEFT({text},1)="\u00b1",'
            f'{"" if upper else "-"}ABS(VALUE(MID({text},2,255))),{pair}),"")')


def make_summary_live(workbook, rows, mean_column, caches):
    sheet = workbook["Summary Report"]
    last_row = 6 + len(rows)
    count_letter = letter(mean_column + 12)
    status_letter = letter(mean_column + 6)
    notes_letter = letter(mean_column + 11)

    def put(row, column, formula, cached=None):
        cell = sheet.cell(row, column)
        previous = caches.get(cell.coordinate, cell.value)
        cell.value = formula
        caches[cell.coordinate] = previous if cached is None else cached

    for r, source in enumerate(rows, 7):
        span = reading_span(r, mean_column)
        invalid = f'COUNTA({span})<>COUNT({span})'
        mean, sd = f'{letter(mean_column)}{r}', f'{letter(mean_column + 1)}{r}'
        for offset, fn in ((0, "AVERAGE"), (4, "MIN"), (5, "MAX")):
            put(r, mean_column + offset, f'=IF({invalid},"Review",IF(COUNT({span})=0,"",{fn}({span})))')
        put(r, mean_column + 1, f'=IF({invalid},"Review",IF(COUNT({span})<2,"N/A",STDEV({span})))')
        for offset, sign in ((2, "+"), (3, "-")):
            put(r, mean_column + offset, f'=IF({invalid},"Review",IF(COUNT({span})<2,"N/A",{mean}{sign}3*{sd}))')
        put(r, mean_column + 12, f'=COUNT({span})', sum(v is not None for v in source.measurements))
        latest = f'LOOKUP(9.99999999999999E+307,{span})'
        put(r, mean_column + 6,
            f'=IFERROR(IF(OR({notes_letter}{r}<>"",{invalid},COUNT({span})=0,'
            f'NOT(ISNUMBER(D{r})),NOT(ISNUMBER(E{r})),D{r}>E{r}),"Review",'
            f'IF(AND({latest}>=D{r},{latest}<=E{r}),"OK","Out of tolerance")),"Review")')
        if not source.explicit_limits:
            for column, upper in ((4, False), (5, True)):
                helper = mean_column + (16 if upper else 15)
                offset = source.upper if upper else source.lower
                cache = float(offset - source.nominal) if offset is not None and source.nominal is not None else ""
                put(r, helper, offset_formula(r, upper), cache)
                helper_cell = f'{letter(helper)}{r}'
                # Unresolved source conflicts must not silently reappear as valid limits.
                conflict = any("conflict" in note.lower() or "unsupported" in note.lower() for note in source.warnings)
                guard = f'{notes_letter}{r}="",' if conflict else ''
                put(r, column, f'=IF(AND({guard}ISNUMBER(C{r}),ISNUMBER({helper_cell})),C{r}+{helper_cell},"")',
                    float(offset) if offset is not None else "")
        for c in [3, 6] + list(range(7, mean_column)):
            font = copy(sheet.cell(r, c).font)
            font.color = "0070C0"
            sheet.cell(r, c).font = font
        sheet.cell(r, 6).comment = Comment(
            "Editable tolerance: use +/-0.5, ±0.5, or +0.5/-0.2. Derived limits update automatically. "
            "If the PDF specifies absolute limits, edit Lower/Upper Tolerance directly instead.", "PDF Intelligence")

    sheet.cell(6, mean_column).comment = Comment(
        "Insert an entire reading/date column immediately before Mean or inside the reading area. "
        "Keep Mean and Tolerance intact. Statistics include the new readings automatically.",
        "PDF Intelligence")
    validation = DataValidation(type="decimal", operator="between", formula1="-1E100", formula2="1E100", allow_blank=True)
    validation.errorTitle = "Enter a numeric measurement"
    validation.error = "Use numbers only. Leave missing measurements blank."
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    sheet.add_data_validation(validation)
    validation.add(f'G7:{letter(mean_column - 1)}{last_row}')

    # One live row per feature; future reading columns update this same table.
    differences = workbook["Differences"]
    header = next(cell.row for cell in differences['A'] if cell.value == 'Measured Feature (Object)')
    prefix = "'Summary Report'!"
    latest_header = header
    differences.cell(header, 5, "Latest Measurement")
    for r, source in enumerate(rows, 7):
        dest_row = latest_header + r - 6
        def latest_put(column, formula, cached):
            cell = differences.cell(dest_row, column, formula)
            cell._style = copy(differences.cell(header + 1, column)._style)
            caches[f'2!{cell.coordinate}'] = cached if cached is not None else ""
        for dest, origin in ((1, 1), (2, 2), (3, mean_column + 7), (4, 3), (7, 6),
                             (8, 4), (9, 5), (10, mean_column + 6), (11, mean_column + 9), (12, mean_column + 8)):
            ref = f'{prefix}{letter(origin)}{r}'
            cell = sheet.cell(r, origin)
            latest_put(dest, f'=IF({ref}="","",{ref})', caches.get(cell.coordinate, cell.value))
        span = reading_span(r, mean_column, prefix)
        latest_value = next(v for v in reversed(source.measurements) if v is not None)
        latest_put(5, f'=IF(COUNTA({span})<>COUNT({span}),"Review",IF(COUNT({span})=0,"",LOOKUP(9.99999999999999E+307,{span})))',
                   float(latest_value))
        latest_put(6, f'=IF(AND(ISNUMBER(D{dest_row}),ISNUMBER(E{dest_row})),E{dest_row}-D{dest_row},"")',
                   float(latest_value - source.nominal) if source.nominal is not None else "")
        # Preserve the exporter height calculated from source labels and pages.
        differences.row_dimensions[dest_row].height = max(
            30, differences.row_dimensions[dest_row].height or 0
        )
    target_row = differences.max_row + 1
    differences.print_area = differences.dimensions
    # Compact reports fit on one page; larger reports retain readable pagination.
    differences.page_setup.fitToHeight = 1 if differences.max_row <= 40 else 0
    for value, color in (("OK", "C6EFCE"), ("Review", "FFEB9C"), ("Out of tolerance", "FFC7CE")):
        differences.conditional_formatting.add(f'J{header + 1}:J{target_row - 1}',
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=PatternFill("solid", fgColor=color)))
    differences["A2"] = "One row per measured feature. Difference = latest reading - nominal. Values update from Summary Report, including new reading columns."
