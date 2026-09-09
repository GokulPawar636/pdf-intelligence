from __future__ import annotations

import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.models.semantic import SemanticDocument
from app.reporting.storage import save_workbook_atomic


def fit_row_heights(sheet, start: int, end: int, max_column: int | None = None) -> None:
    """Give wrapped cells enough room while keeping very long evidence manageable."""
    for row in sheet.iter_rows(min_row=start, max_row=end, max_col=max_column):
        lines = 1
        for cell in row:
            if cell.value is None or cell.data_type == "f":
                continue
            width = max(8, sheet.column_dimensions[cell.column_letter].width - 2)
            lines = max(lines, sum(max(1, math.ceil(len(part) / width))
                                   for part in str(cell.value).split("\n")))
        sheet.row_dimensions[row[0].row].height = min(150, max(28, lines * 15 + 10))


def _display_field_name(name: str) -> str:
    return re.sub(r"\s+", " ", name.replace("_", " ")).strip().title()


def _text(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def extract_title_fields(document: dict) -> list[tuple[str, str]]:
    """Extract leading label/value pairs without assuming a document template."""
    title_fields: list[tuple[str, str]] = []
    seen: set[str] = set()
    for record in document.get("records", []):
        tokens = [
            str(field.get("value", {}).get("raw_value") or "").strip()
            for field in record.get("fields", [])
        ]
        for index, token in enumerate(tokens):
            if not token:
                continue
            if token.endswith(":") and index + 1 < len(tokens) and tokens[index + 1]:
                label, value = token[:-1].strip(), tokens[index + 1]
            elif (index + 2 < len(tokens) and tokens[index + 1].endswith(":")
                  and tokens[index + 2]):
                label, value = f"{token} {tokens[index + 1][:-1].strip()}", tokens[index + 2]
            else:
                continue
            key = re.sub(r"\s+", " ", label).strip().casefold()
            if any(existing.endswith(" " + key) for existing in seen):
                continue
            shorter = {existing for existing in seen if key.endswith(" " + existing)}
            if shorter:
                title_fields[:] = [(label, value) for label, value in title_fields
                                    if re.sub(r"\s+", " ", label).strip().casefold() not in shorter]
                seen.difference_update(shorter)
            if key and key not in seen:
                title_fields.append((label, value))
                seen.add(key)
    return title_fields


def _append(sheet, values: list[Any]) -> None:
    """PDF strings are literal data, including strings beginning with '='."""
    row = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        value = _text(value)
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            cell.data_type = "s"
            if len(value) > 32767:
                cell.comment = Comment("Long value: see all parts in the Evidence sheet.", "PDF Intelligence")


def _style(sheet, table_index: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cells in sheet.iter_cols():
        length = max((max((len(line) for line in str(c.value or "").splitlines()), default=0)
                      for c in cells[:101]), default=12)
        sheet.column_dimensions[cells[0].column_letter].width = min(55, max(14, length + 2))
    fit_row_heights(sheet, 2, sheet.max_row)
    sheet.sheet_view.zoomScale = 90
    sheet.print_title_rows = "1:1"
    sheet.print_area = sheet.dimensions
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    if sheet.max_row > 1:
        table = Table(displayName=f"DocumentTable{table_index}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def _sheet_name(workbook: Workbook, title: str) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", ILLEGAL_CHARACTERS_RE.sub("", title)).strip().strip("'")[:31] or "Records"
    if base.casefold() == "history":
        base = "History Records"
    name, suffix = base, 2
    while name.casefold() in {s.casefold() for s in workbook.sheetnames}:
        tail = f" {suffix}"
        name = base[:31 - len(tail)] + tail
        suffix += 1
    return name


def _slots(record: dict) -> list[tuple[str, int]]:
    counts: dict[str, int] = {}
    slots = []
    for field in record.get("fields", []):
        name = field["name"]
        counts[name] = counts.get(name, 0) + 1
        slots.append((name, counts[name]))
    return slots


def export_semantic_to_excel(input_path: str | Path, output_path: str | Path) -> Path:
    """Export input-driven record groups plus a complete field-level audit trail."""
    input_file = Path(input_path).expanduser().resolve()
    output_file = Path(output_path).expanduser().resolve()
    document = SemanticDocument.model_validate_json(input_file.read_text(encoding="utf-8")).model_dump(mode="json")
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    evidence = workbook.create_sheet("Evidence")
    _append(overview, ["Property", "Value"])
    metadata = document.get("document_metadata", {})
    for name, value in [
        ("Source PDF", metadata.get("file_name", input_file.name)),
        ("Pages", metadata.get("page_count")),
        ("Records", len(document["records"])),
        ("Extraction method", document["extraction_provider"]),
        ("Extraction confidence", document["extraction_confidence"]),
        ("Reading this workbook", "Each data sheet groups related records. Columns come from the document. Evidence preserves raw values, units, and source locations."),
    ]:
        _append(overview, [name, value])
    for note in document["global_notes"] + document["warnings"]:
        _append(overview, ["Note", note])
    if not document["records"]:
        _append(overview, ["Extraction status", "No records extracted. Check whether the PDF contains readable text; scanned pages may require OCR."])
    _append(evidence, ["Sheet", "Record ID", "Record Type", "Field", "Occurrence", "Part",
                       "Raw Value", "Normalized Value", "Unit", "Page", "Source ID",
                       "Source Text", "Bounding Box", "Method", "Confidence", "Warnings"])

    groups: dict[tuple, list[dict]] = {}
    for record in document["records"]:
        # Generic tables with different schemas must not become one sparse table.
        # Headerless tables are kept separate because positional names convey no meaning.
        names = tuple(sorted({f["name"] for f in record["fields"]}))
        table_id = ""
        if record["record_type"] == "table_row" and all(re.fullmatch(r"column_\d+", n) for n in names):
            table_id = (record.get("record_id") or "").rsplit("_r", 1)[0]
        key = (record["record_type"], names if record["record_type"] == "table_row" else (), table_id)
        groups.setdefault(key, []).append(record)

    for (record_type, _, table_id), records in groups.items():
        slots = list(dict.fromkeys(slot for record in records for slot in _slots(record)))
        title = _display_field_name(record_type)
        if record_type == "table_row":
            title = table_id or " & ".join(_display_field_name(name) for name, _ in slots[:2])
        sheet = workbook.create_sheet(_sheet_name(workbook, title))
        headers, used = [], {"record id", "source pages", "confidence", "warnings"}
        for name, occurrence in slots:
            base = _display_field_name(name) or "Value"
            if occurrence > 1:
                base += f" ({occurrence})"
            label, suffix = base, 2
            while label.casefold() in used:
                label = f"{base} ({suffix})"
                suffix += 1
            used.add(label.casefold())
            headers.append(label)
        _append(sheet, headers + ["Record ID", "Source Pages", "Confidence", "Warnings"])
        _append(overview, [f"Sheet: {sheet.title}", f"{len(records)} records; {len(slots)} document fields"])
        for record in records:
            values = {}
            for slot, field in zip(_slots(record), record["fields"]):
                value = field["value"]
                normalized = value.get("normalized_value")
                values[slot] = normalized if normalized is not None else value.get("raw_value")
                source = value.get("source") or {}
                raw = str(value.get("raw_value") or "")
                normalized_text = _text(normalized)
                source_text = source.get("source_text") or ""
                strings = [raw, str(normalized_text) if normalized_text is not None else "", source_text]
                parts = max(1, *( (len(s) + 29999) // 30000 for s in strings))
                for part in range(parts):
                    chunks = [s[part * 30000:(part + 1) * 30000] for s in strings]
                    _append(evidence, [sheet.title, record.get("record_id"), record_type,
                        field["name"], slot[1], part + 1, chunks[0],
                        normalized if parts == 1 and isinstance(normalized, (int, float, bool)) else chunks[1],
                        value.get("unit"), source.get("page_number"), source.get("source_id"),
                        chunks[2], source.get("bbox"), source.get("extraction_method"),
                        source.get("confidence"), "\n".join(record.get("warnings", []))])
            _append(sheet, [values.get(slot) for slot in slots] + [record.get("record_id"),
                ", ".join(map(str, record["source_pages"])), record["confidence"],
                "\n".join(record.get("warnings", []))])
            for index, slot in enumerate(slots, 1):
                field = next((f for s, f in zip(_slots(record), record["fields"]) if s == slot), None)
                if field:
                    value = field["value"]
                    cell = sheet.cell(sheet.max_row, index)
                    raw = value.get("raw_value") or ""
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        match = re.fullmatch(r"[+-]?\d+\.(\d+)", raw)
                        if match:
                            cell.number_format = "0." + "0" * min(15, len(match[1]))
                    if value.get("unit"):
                        cell.comment = Comment(f"Unit: {value['unit']}\nRaw value: {raw[:1000]}", "PDF Intelligence")
    # Keep the audit trail last, with the document-driven data sheets together.
    workbook.move_sheet(evidence, offset=len(workbook.worksheets) - 2)
    for index, sheet in enumerate(workbook.worksheets, 1):
        _style(sheet, index)
    workbook.active = 0
    output_file.parent.mkdir(parents=True, exist_ok=True)
    return save_workbook_atomic(workbook, output_file)


