"""Validate an XLSX before atomically publishing it to its destination."""
from __future__ import annotations

import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZipFile

from openpyxl import load_workbook


def save_workbook_atomic(workbook, destination, finalize=None) -> Path:
    destination = Path(destination).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=destination.parent, suffix=".xlsx", prefix=".building-", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        if finalize:
            finalize(temporary)
        with ZipFile(temporary) as archive:
            if archive.testzip() is not None:
                raise ValueError("Workbook archive validation failed.")
        check = load_workbook(temporary, read_only=True, data_only=True)
        try:
            for sheet in check:
                if sheet.max_column > 16384 or sheet.max_row > 1048576:
                    raise ValueError("Excel worksheet size limit exceeded.")
                for row in sheet:
                    for cell in row:
                        if cell.data_type == "e":
                            raise ValueError(f"Excel error in {sheet.title}!{cell.coordinate}: {cell.value}")
        finally:
            check.close()
        try:
            os.replace(temporary, destination)
        except PermissionError as exc:
            raise PermissionError(
                f"Cannot publish {destination.name}; it may be open in Excel. "
                "Choose another output path or close that workbook. The previous file was preserved."
            ) from exc
        return destination
    finally:
        temporary.unlink(missing_ok=True)
