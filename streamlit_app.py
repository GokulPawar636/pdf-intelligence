"""Multiple PDF uploads with a side-by-side, formatted Excel preview."""
from __future__ import annotations

from html import escape
from base64 import b64encode
from io import BytesIO
import hashlib
import json
import logging
import os
import re
from pathlib import Path
import tempfile

import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.batch import run_batch
from app.reporting.excel_importer import import_excel


def clear_result():
    for key in ("result", "attempted", "preview_sheet", "conversion_error"):
        st.session_state.pop(key, None)


def convert_uploads(files, progress=None, *, strict=False, use_ai=False, mappings=None):
    if not files:
        raise ValueError("Choose at least one PDF.")
    if any(Path(name).suffix.lower() in ('.xlsx', '.xls') for name, _ in files):
        if len(files) != 1:
            raise ValueError('Upload one Excel workbook at a time, separately from PDFs.')
        name, contents = files[0]
        if len(contents) > 100*1024*1024:
            raise ValueError('Excel input exceeds 100 MB.')
        with tempfile.TemporaryDirectory(prefix='excel-import-') as folder:
            source = Path(folder) / Path(name.replace('\\', '/')).name
            source.write_bytes(contents)
            return import_excel(source, Path(folder)/'Converted_Report.xlsx', strict=strict, use_ai=use_ai, mappings=mappings)
    if sum(len(contents) for _, contents in files) > 500 * 1024 * 1024:
        raise ValueError("The selected PDFs exceed 500 MB. Please use a smaller batch.")
    with tempfile.TemporaryDirectory(prefix="pdf-intelligence-") as workspace:
        paths = []
        for index, (name, contents) in enumerate(files):
            safe_name = Path(name.replace("\\", "/")).name
            if not contents or Path(safe_name).suffix.lower() != ".pdf":
                raise ValueError(f"{safe_name or 'File'}: choose a non-empty PDF.")
            if len(contents) > 100 * 1024 * 1024:
                raise ValueError(f"{safe_name}: the per-file limit is 100 MB.")
            folder = Path(workspace) / "inputs" / str(index)
            folder.mkdir(parents=True)
            source = folder / safe_name
            source.write_bytes(contents)
            paths.append(source)
        output = run_batch(paths, Path(workspace) / "output", progress=progress, strict=strict, offline=not use_ai)
        manifest = json.loads(output.with_name("manifest.json").read_text(encoding="utf-8"))
        return {"name": output.name, "contents": output.read_bytes(),
                "warnings": manifest["warnings"], "source": ", ".join(name for name, _ in files),
                "files": len(files), "pages": manifest["pages"],
                "measurement_columns": manifest["measurement_columns"]}


def convert_upload(name, contents):
    return convert_uploads([(name, contents)])


def preview_html(contents, selected):
    workbook = load_workbook(BytesIO(contents), data_only=True)
    try:
        sheet = workbook[selected]
        pictures = {}
        for picture in sheet._images:
            anchor = picture.anchor._from
            pictures[(anchor.row+1, anchor.col+1)] = (
                '<img alt="Company logo" style="width:300px;max-width:100%;height:auto" src="data:image/'
                + picture.format + ';base64,' + b64encode(picture._data()).decode('ascii') + '">')
        hidden = set()
        for dimension in sheet.column_dimensions.values():
            if dimension.hidden:
                hidden.update(range(dimension.min, dimension.max + 1))
        columns = [c for c in range(1, sheet.max_column + 1) if c not in hidden][:80]
        max_row = min(sheet.max_row, 200)
        merges, covered = {}, set()
        for merged in sheet.merged_cells.ranges:
            visible = [c for c in columns if merged.min_col <= c <= merged.max_col]
            if not visible or merged.min_row > max_row:
                continue
            merges[(merged.min_row, visible[0])] = (len(visible), min(merged.max_row, max_row) - merged.min_row + 1)
            for r in range(merged.min_row, min(merged.max_row, max_row) + 1):
                for c in visible:
                    if (r, c) != (merged.min_row, visible[0]):
                        covered.add((r, c))
        body = []
        for r in range(1, max_row + 1):
            cells = []
            for c in columns:
                if (r, c) in covered:
                    continue
                cell = sheet.cell(r, c)
                value = cell.value
                if isinstance(value, (int, float)) and re.fullmatch(r'0\.0{1,15}', cell.number_format):
                    precision = len(cell.number_format)-2
                    text = f"{value:.{precision}f}"
                else:
                    text = "" if value is None else str(value)
                colspan, rowspan = merges.get((r, c), (1, 1))
                color = cell.fill.fgColor.rgb if cell.fill.fgColor.type == "rgb" else None
                fg = cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None
                styles = ["padding:9px", "border:1px solid #d4dce5", "min-width:95px",
                          "max-width:450px", "overflow-wrap:anywhere", "vertical-align:middle", "white-space:pre-wrap"]
                if cell.fill.patternType == "solid" and isinstance(color, str):
                    styles.append("background:#" + color[-6:])
                if isinstance(fg, str):
                    styles.append("color:#" + fg[-6:])
                if cell.font.bold:
                    styles.append("font-weight:600")
                if cell.alignment.horizontal in ("left", "right", "center"):
                    styles.append("text-align:" + cell.alignment.horizontal)
                if value in ("OK", "Review", "Out of tolerance"):
                    styles.append("background:" + {"OK": "#e2f3e8", "Review": "#fff0bc", "Out of tolerance": "#ffe1e1"}[value])
                    styles.append("color:#172638")
                content = pictures.get((r,c), '') + escape(text)
                cells.append(f'<td colspan="{colspan}" rowspan="{rowspan}" style="{";".join(styles)}">{content}</td>')
            body.append("<tr>" + "".join(cells) + "</tr>")
        return ('<div style="overflow:auto;max-height:680px;background:white;color:#172638">'
                '<table style="border-collapse:collapse;font:13px Arial,sans-serif">' +
                "".join(body) + "</table></div>")
    finally:
        workbook.close()


def main():
    st.set_page_config(page_title="PDF to Excel", page_icon=":page_facing_up:", layout="wide")
    st.title("Inspection reports to Excel")
    st.caption("Convert PDF or Excel measurements into one clear report.")
    upload_panel, preview_panel = st.columns([1, 2.8], gap="large")
    with upload_panel:
        input_type = st.radio('Input format', ['PDF', 'Excel'], horizontal=True, on_change=clear_result)
        st.subheader("Upload " + input_type)
        use_ai = os.getenv('GROQ_ENABLED', 'false').lower() in ('true', '1', 'yes')
        uploaded = st.file_uploader("Choose one or more PDFs" if input_type == 'PDF' else 'Choose an Excel workbook',
                                    type=['pdf'] if input_type == 'PDF' else ['xlsx','xls'],
                                    accept_multiple_files=input_type == 'PDF', key='upload_'+input_type, on_change=clear_result)
        if uploaded and not isinstance(uploaded, list):
            uploaded = [uploaded]
        st.caption('Upload PDFs together.' if input_type == 'PDF' else 'Upload one .xlsx or .xls workbook.')
        if use_ai:
            st.caption('AI assistance is enabled. Unfamiliar content may be sent to Groq.')
    with preview_panel:
        st.subheader("Excel preview")
        if uploaded:
            files = [(item.name, item.getvalue()) for item in uploaded]
            fingerprint = hashlib.sha256(f"simple-v5:{input_type}:{use_ai}".encode())
            for name, contents in files:
                fingerprint.update(name.encode())
                fingerprint.update(hashlib.sha256(contents).digest())
            token = fingerprint.hexdigest()
            if st.session_state.get("attempted") != token:
                st.session_state["attempted"] = token
                st.session_state.pop("result", None)
                st.session_state.pop("conversion_error", None)
                bar = st.progress(0, text="Preparing report...")
                try:
                    def progress(done, total, name):
                        bar.progress(done / total, text=f"{done}/{total} files processed - {name}")
                    st.session_state["result"] = convert_uploads(files, progress, strict=False, use_ai=use_ai)
                except ValueError as exc:
                    st.session_state["conversion_error"] = str(exc)
                except Exception:
                    logging.getLogger(__name__).exception("Batch conversion failed")
                    st.session_state["conversion_error"] = "The report could not be generated. Check your files and retry."
                finally:
                    bar.empty()
        result = st.session_state.get("result")
        if st.session_state.get("conversion_error"):
            st.error(st.session_state["conversion_error"])
            st.button('Retry', on_click=clear_result)
        if result is None:
            st.info("Upload files on the left. The Excel preview will appear here automatically after processing.")
            return
        if result["warnings"]:
            st.warning("Review report ready. Image-only content and uncertain interpretations may be incomplete; check the notes against your sources before using results.")
            with st.expander("Review notes"):
                for note in result["warnings"]:
                    st.write(note)
        else:
            st.success("Your Excel report is ready. Automated checks found no flagged issues.")
        st.download_button("Download Excel", result["contents"], file_name=result["name"],
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
        workbook = load_workbook(BytesIO(result["contents"]), read_only=True)
        names = workbook.sheetnames
        workbook.close()
        if st.session_state.get("preview_sheet") not in names:
            st.session_state["preview_sheet"] = names[0]
        selected = st.selectbox("Worksheet", names, key="preview_sheet")
        st.html(preview_html(result["contents"], selected))
        st.caption("Preview: up to 200 rows and 80 visible columns. Download includes the complete workbook with live formulas.")


if __name__ == "__main__":
    main()
