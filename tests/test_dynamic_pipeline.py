import json

import fitz
import pytest
from openpyxl import load_workbook

from app.semantic.dynamic import normalize_value
from app.semantic.extractor import SemanticExtractor
from pipeline_main import run_pipeline


@pytest.mark.parametrize("value,name,expected", [
    ("00123", "code", "00123"), ("12345678901234567", "value", "12345678901234567"),
    ("001.50", "value", "001.50"), ("123", "account_number", "123"),
    ("1,234", "value", "1,234"), ("12.50", "amount", 12.5), ("0", "count", 0),
])
def test_normalization_preserves_ambiguous_values(value, name, expected):
    assert normalize_value(value, name) == expected


def table(rows, table_id="p001_t001"):
    return {"table_id": table_id, "confidence": 0.85, "cells": [
        {"row": r, "column": c, "text": text}
        for r, values in enumerate(rows) for c, text in enumerate(values)]}


def test_numeric_first_row_and_text_only_tables_are_not_dropped():
    result = SemanticExtractor(offline=True).extract({"pages": [{"page_number": 1, "tables": [
        table([["001", "42"], ["002", "50"]]),
        table([["Asha", "Sales"], ["Lee", "Support"]], "p001_t002"),
    ]}]})
    assert len(result.records) == 4
    assert result.records[0].fields[0].value.raw_value == "001"
    assert result.records[2].fields[0].value.raw_value == "Asha"
    assert all(r.warnings for r in result.records)
    assert result.extraction_provider == "local"


def test_duplicate_table_rows_survive_online_merge():
    result = SemanticExtractor(api_key="test-key").extract({"pages": [{"page_number": 1, "tables": [
        table([["Product", "Amount"], ["Widget", "10"], ["Widget", "10"]]),
    ]}]})
    assert len(result.records) == 2
    assert len({r.record_id for r in result.records}) == 2


def test_unicode_headers_and_duplicate_headers():
    result = SemanticExtractor(offline=True).extract({"pages": [{"page_number": 1, "tables": [
        table([["नाम", "Amount", "Amount"], ["Asha", "10", "20"]]),
    ]}]})
    names = [f.name for f in result.records[0].fields]
    assert names[0] != "column_1"
    assert names[1:] == ["amount", "amount_2"]


def test_form_labels_are_not_mistaken_for_column_headers():
    result = SemanticExtractor(offline=True).extract({"pages": [{"page_number": 1, "tables": [
        table([["Customer:", "Asha"], ["Account:", "123"]]),
    ]}]})
    assert len(result.records) == 2
    assert result.records[0].fields[0].value.raw_value == "Customer:"


def test_unassigned_words_inside_table_bounds_remain_available():
    bbox = {"x0": 0, "y0": 0, "x1": 100, "y1": 100}
    region = {"region_type": "text_region", "text": "A note", "bbox": bbox, "element_indexes": [3]}
    regions = SemanticExtractor._text_regions({"regions": [region], "tables": [
        {"bbox": bbox, "cells": [{"source_word_indexes": [1, 2]}]},
    ]})
    assert regions == [region]


@pytest.mark.parametrize("kind", ["invoice", "inventory", "prose", "blank"])
def test_pdf_to_dynamic_workbook_end_to_end(tmp_path, kind):
    path = tmp_path / f"{kind}.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    if kind in ("invoice", "inventory"):
        rows = ([["Item", "Amount"], ["Widget", "42.50"], ["Service", "100.00"]]
                if kind == "invoice" else [["Product", "Stock"], ["Bolt", "12"], ["Nut", "25"]])
        page.insert_text((50, 40), f"{kind.title()} document")
        for r, values in enumerate(rows):
            for c, value in enumerate(values):
                page.insert_text((50 + c * 180, 100 + r * 25), value)
    elif kind == "prose":
        page.insert_text((50, 50), "Project: Aurora")
        page.insert_text((50, 100), "The team completed the first review and recorded its findings.")
    pdf.save(path)
    pdf.close()
    output = run_pipeline(path, tmp_path / "outputs", offline=True)
    workbook = load_workbook(output)
    semantic = json.loads((output.parent / "intermediate/semantic_document.json").read_text(encoding="utf-8"))
    assert output.name == f"{kind}.xlsx"
    assert "Inspection Results" not in workbook.sheetnames
    values = [f["value"]["raw_value"] for r in semantic["records"] for f in r["fields"]]
    if kind in ("invoice", "inventory"):
        assert rows[1][0] in values and rows[1][1] in values
        table_records = [r for r in semantic["records"] if r["record_type"] == "table_row"]
        assert len(table_records) == 2
        assert rows[0][1].lower() in {f["name"] for r in table_records for f in r["fields"]}
    elif kind == "prose":
        assert "Aurora" in values
        assert any("completed the first review" in v for v in values)
    else:
        assert not semantic["records"]
        assert semantic["warnings"]
