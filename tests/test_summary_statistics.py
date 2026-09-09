import json

from openpyxl import load_workbook
import pytest

from app.reporting.summary_exporter import export_summary
from app.reporting.summary_exporter import build_summary_rows


def test_statistics_and_ambiguous_limits(tmp_path):
    fields = [{"name": name, "value": {"raw_value": raw}} for name, raw in [
        ("object", "=Literal feature"), ("control", "Diameter"), ("nominal", "10"),
        ("tolerance", "0.5"), ("measured", "9"), ("measured", "11"),
    ]]
    semantic = {"source_schema": "test", "source_schema_version": "1", "document_id": "test",
                "extraction_model": "local", "records": [{"record_type": "measurement", "fields": fields}]}
    structured = {"source_schema": "test", "source_schema_version": "1", "document": {}, "pages": []}
    source = tmp_path / "semantic.json"
    layout = tmp_path / "structured.json"
    source.write_text(json.dumps(semantic), encoding="utf-8")
    layout.write_text(json.dumps(structured), encoding="utf-8")
    path = export_summary(source, layout, tmp_path / "summary.xlsx")
    values = load_workbook(path, data_only=True).active
    formulas = load_workbook(path).active
    assert values["A7"].value == "=Literal feature"
    assert formulas["A7"].data_type == "s"
    assert values["D7"].value is None and values["E7"].value is None
    assert values["I7"].value == 10
    assert values["J7"].value == pytest.approx(2**0.5)
    assert values["K7"].value == pytest.approx(10 + 3 * 2**0.5)
    assert values["L7"].value == pytest.approx(10 - 3 * 2**0.5)
    assert values["M7"].value == 9 and values["N7"].value == 11
    assert values["O7"].value == "Review"
    assert formulas["I7"].value == '=IF(COUNT(G7:H7)=0,"",AVERAGE(G7:H7))'


def test_numbered_readings_and_safe_grouping():
    def record(feature, unit, nominal, readings):
        values = [("Object", feature), ("Control", "Diameter"), ("Nominal", nominal),
                  ("Tolerance", "+/-0.5"), ("Unit", unit)] + readings
        return {"fields": [{"name": k, "value": {"raw_value": v}} for k, v in values]}
    rows = build_summary_rows({"records": [
        record("A", "mm", "10", [("Reading 1", "9.9"), ("Reading 2", "10.1")]),
        record("A", "mm", "10", [("Actual Value", "10.2")]),
        record("B", "mm", "10", [("Measured", "10")]),
        record("A", "inch", "10", [("Measured", "10")]),
        record("A", "mm", "20", [("Measured", "20")]),
    ]}, {})
    assert len(rows) == 4
    assert list(map(float, rows[0].measurements)) == [9.9, 10.1, 10.2]
    assert all(len(r.measurements) == 1 for r in rows[1:])


def test_different_pdf_layout_automatically_calculates_statistics(tmp_path):
    import fitz
    from pipeline_main import run_pipeline

    pdf = fitz.open()
    page = pdf.new_page()
    data = [["Feature", "Control", "Nominal", "Reading 1", "Reading 2", "Unit"],
            ["Shaft", "Diameter", "10", "9", "11", "mm"],
            ["Pin", "Diameter", "4", "3.9", "4.1", "mm"]]
    for r, values in enumerate(data):
        for c, value in enumerate(values):
            page.insert_text((40 + c * 85, 70 + r * 25), value, fontsize=9)
    source = tmp_path / "different_layout.pdf"
    pdf.save(source)
    pdf.close()
    output = run_pipeline(source, tmp_path / "output", offline=True)
    sheet = load_workbook(output, data_only=True).active
    assert sheet.title == "Summary Report"
    assert sheet["A7"].value == "Shaft"
    assert sheet["I7"].value == 10
    assert sheet["J7"].value == pytest.approx(2 ** 0.5)
    assert sheet["K7"].value == pytest.approx(10 + 3 * 2 ** 0.5)
    assert sheet["L7"].value == pytest.approx(10 - 3 * 2 ** 0.5)
