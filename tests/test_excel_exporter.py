import json

from openpyxl import load_workbook
from app.reporting.excel_exporter import export_semantic_to_excel


def field(name, raw, normalized=None, unit=None):
    return {"name": name, "value": {"raw_value": raw, "normalized_value": normalized, "unit": unit}}


def export(tmp_path, records):
    payload = {"source_schema": "test", "source_schema_version": "1", "document_id": "test",
               "extraction_model": "deterministic", "records": records}
    source = tmp_path / "input.json"
    source.write_text(json.dumps(payload), encoding="utf-8")
    return load_workbook(export_semantic_to_excel(source, tmp_path / "result.xlsx"))


def test_workbook_groups_content_without_inspection_template(tmp_path):
    workbook = export(tmp_path, [
        {"record_type": "invoice_item", "fields": [field("description", "Widget"), field("amount", "42.50", 42.5)]},
        {"record_type": "employee", "fields": [field("name", "Asha"), field("department", "Sales")]},
    ])
    assert workbook.sheetnames == ["Overview", "Invoice Item", "Employee", "Evidence"]
    assert workbook["Invoice Item"]["B2"].value == 42.5
    assert workbook["Invoice Item"]["B2"].number_format == "0.00"
    assert "Department" not in [c.value for c in workbook["Invoice Item"][1]]
    assert workbook["Evidence"].max_row == 5


def test_duplicate_fields_nested_values_units_and_literal_formulas_survive(tmp_path):
    workbook = export(tmp_path, [{"record_type": "details", "fields": [
        field("tag", "first"), field("tag", "second"), field("formula", "=1+1"),
        field("options", "source", {"enabled": True}), field("distance", "2.50", 2.5, "km"),
    ]}])
    sheet = workbook["Details"]
    assert sheet["A2"].value == "first"
    assert sheet["B2"].value == "second"
    assert sheet["C2"].value == "=1+1" and sheet["C2"].data_type == "s"
    assert json.loads(sheet["D2"].value) == {"enabled": True}
    assert "km" in sheet["E2"].comment.text
    assert all(c.data_type != "f" for s in workbook for row in s for c in row)


def test_unrelated_table_schemas_and_sheet_name_collisions(tmp_path):
    workbook = export(tmp_path, [
        {"record_type": "table_row", "fields": [field("product", "Widget"), field("price", "5", 5)]},
        {"record_type": "table_row", "fields": [field("person", "Lee"), field("role", "Editor")]},
        {"record_type": "overview", "fields": [field("record_id", "001")]},
        {"record_type": "bad/title:*?[]" * 8, "fields": [field("name", "test")]},
    ])
    assert "Product & Price" in workbook.sheetnames
    assert "Person & Role" in workbook.sheetnames
    assert "Overview 2" in workbook.sheetnames
    assert all(len(s) <= 31 for s in workbook.sheetnames)
    assert len({s.casefold() for s in workbook.sheetnames}) == len(workbook.sheetnames)
    for sheet in workbook:
        headers = [c.value for c in sheet[1]]
        assert len(headers) == len(set(str(h).casefold() for h in headers))


def test_long_text_is_complete_in_evidence(tmp_path):
    raw = "Long passage " * 4000
    workbook = export(tmp_path, [{"record_type": "document_text", "fields": [field("text", raw)]}])
    evidence = workbook["Evidence"]
    assert "".join(evidence.cell(r, 7).value or "" for r in range(2, evidence.max_row + 1)) == raw
    assert workbook["Document Text"]["A2"].comment is not None


def test_empty_document_has_clear_status(tmp_path):
    workbook = export(tmp_path, [])
    assert workbook.sheetnames == ["Overview", "Evidence"]
    assert any("No records extracted" in str(c.value) for row in workbook["Overview"] for c in row)
