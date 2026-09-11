from decimal import Decimal
import json
from pathlib import Path

import fitz
from openpyxl import load_workbook
import pytest

from app.batch import combine_rows, run_batch
from app.reporting.summary_exporter import SummaryRow


def source(name, entries, part="P1"):
    return {"name": name, "details": {"part": part, "date": "9/11/2026", "serial": name,
             "unit": "mm", "alignment": "world"},
            "rows": [SummaryRow(feature, "Diameter", Decimal("10"), Decimal("9.5"), Decimal("10.5"),
                                 "+/-0.5", [Decimal(value)], unit="mm") for feature, value in entries]}


def test_matching_uses_identity_not_row_order_and_preserves_gaps():
    rows, headers = combine_rows([source("one", [("A", "10"), ("B", "9.8")]),
                                  source("two", [("B", "10.1")]),
                                  source("three", [("B", "10.2"), ("A", "9.9")])])
    assert len(headers) == 3 and len(rows) == 2
    assert rows[0].measurements == [Decimal("10"), None, Decimal("9.9")]
    assert rows[1].measurements == [Decimal("9.8"), Decimal("10.1"), Decimal("10.2")]


def test_different_parts_or_specifications_never_pool_readings():
    a, b = source("one", [("A", "10")]), source("two", [("A", "10.1")], part="P2")
    rows, _ = combine_rows([a, b])
    assert len(rows) == 2
    b["details"]["part"] = "P1"
    b["rows"][0].upper = Decimal("11")
    assert len(combine_rows([a, b])[0]) == 2


def make_pdf(path, value, pages=1):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((40, 40), "Part Number: P1")
    page.insert_text((40, 60), "Date: 9/11/2026")
    for r, row in enumerate([["Object", "Control", "Nominal", "Tolerance", "Reading", "Unit"],
                              ["Shaft", "Diameter", "10", "+/-0.5", str(value), "mm"]]):
        for c, text in enumerate(row):
            page.insert_text((40+c*85, 110+r*25), text, fontsize=9)
    for n in range(1, pages):
        doc.new_page().insert_text((40, 40), f"Additional page {n}")
    doc.save(path)
    doc.close()
    return path


@pytest.mark.parametrize("count", [1, 2, 3])
def test_variable_file_and_page_counts_create_one_workbook(tmp_path, count):
    inputs = [make_pdf(tmp_path / f"sample{i}.pdf", 9.9+i*.1, pages=i+1) for i in range(count)]
    path = run_batch(inputs, tmp_path / "reports")
    book = load_workbook(path, data_only=True)
    try:
        assert book.sheetnames == ["Summary Report", "Differences"]
        assert book.active["A7"].value == "Shaft"
        assert book.active.cell(7, 7+count).value == pytest.approx(9.9+(count-1)*.05)
        assert book.active.cell(7, 8+count).value == ("N/A" if count == 1 else pytest.approx(.1/(2**.5) if count == 2 else .1))
        assert book.active.cell(5, 6+count).value == f"PDF {count}"
    finally:
        book.close()
    manifest = json.loads(path.with_name("manifest.json").read_text())
    assert manifest["pages"] == sum(range(1, count+1))
    assert manifest["measurement_columns"] == count
    assert len(list((tmp_path / "reports").rglob("*.xlsx"))) == 1


def test_duplicate_pdf_is_not_counted_twice(tmp_path):
    pdf = make_pdf(tmp_path / "one.pdf", 10)
    output = run_batch([pdf, pdf], tmp_path / "out")
    manifest = json.loads(output.with_name("manifest.json").read_text())
    assert manifest["measurement_columns"] == 1
    assert manifest["inputs"][1]["duplicate"] is True


def test_known_measurements_skip_groq_even_when_enabled(tmp_path, monkeypatch):
    import app.batch as batch
    original = batch.extract_file
    calls = []
    def local_only(*args, **kwargs):
        calls.append(kwargs.get('offline'))
        assert kwargs.get('offline') is True
        return original(*args, **kwargs)
    monkeypatch.setattr(batch, 'extract_file', local_only)
    output = run_batch([make_pdf(tmp_path/'known.pdf', 10)], tmp_path/'out', offline=False)
    assert output.exists() and calls == [True]
    manifest = json.loads(output.with_name('manifest.json').read_text())
    assert manifest['inputs'][0]['extraction_strategy'] == 'local'


def test_invalid_batch_does_not_publish_partial_workbook(tmp_path):
    pdf = make_pdf(tmp_path / "one.pdf", 10)
    invalid = tmp_path / "invalid.pdf"
    invalid.write_bytes(b"not a pdf")
    with pytest.raises(Exception):
        run_batch([pdf, invalid], tmp_path / "out")
    assert not list((tmp_path / "out").rglob("*.xlsx"))


def test_general_document_is_consolidated_without_fragmented_tabs(tmp_path):
    pdf = fitz.open()
    pdf.new_page().insert_text((40, 40), "Project: Aurora")
    path = tmp_path / "general.pdf"
    pdf.save(path)
    pdf.close()
    output = run_batch([path], tmp_path / "out")
    book = load_workbook(output, data_only=True)
    try:
        assert book.sheetnames == ["Summary Report", "Differences"]
        assert any(row[4] == "Aurora" for row in book.active.iter_rows(min_row=2, values_only=True))
    finally:
        book.close()
