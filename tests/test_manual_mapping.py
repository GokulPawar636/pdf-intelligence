import pytest
from openpyxl import Workbook, load_workbook
from app.reporting.excel_importer import import_excel, numeric


def test_custom_headers_map_without_ai_or_renaming_source(tmp_path):
    source = tmp_path/'custom.xlsx'
    book = Workbook()
    book.active.append(['Unusual report title'])
    book.active.append(['Item tag','Test kind','Design','Allowance','Scale','Run A','Run B'])
    book.active.append(['Shaft','Diameter',1,'+/-0.5','mm',1e-6,1.1])
    book.save(source)
    mapping = {'Sheet': {'header_row':1,'columns':{'object':0,'control':1,'nominal':2,
                'tolerance':3,'unit':4,'lower':None,'upper':None},'readings':[5,6]}}
    result = import_excel(source,tmp_path/'out.xlsx',mappings=mapping,strict=True)
    assert result['measurement_columns'] == 2
    output = load_workbook(tmp_path/'out.xlsx',data_only=True)
    assert output.active['G7'].value == 1e-6
    assert output.active['H7'].value == 1.1
    assert output.active['I7'].value == pytest.approx(.5500005)
    output.close()
    mapping['Sheet']['readings'] = [2]
    with pytest.raises(ValueError,match='overlap'):
        import_excel(source,tmp_path/'invalid.xlsx',mappings=mapping)
    assert not (tmp_path/'invalid.xlsx').exists()


@pytest.mark.parametrize('value',[True,'NaN','Infinity','1e101','1.1234567890123456'])
def test_unsupported_excel_numbers_are_rejected(value):
    assert numeric(value) is None
