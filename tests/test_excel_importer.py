from openpyxl import Workbook, load_workbook
import pytest
from app.reporting.excel_importer import import_excel


def source(path, reading=1.0001):
    book = Workbook()
    book.active.append(['Report title'])
    book.active.append(['Units','Object','Control','Nominal','Tolerance','Sample 1','Sample 2'])
    book.active.append(['mm','Shaft','Diameter',1,'+/-0.1',reading,1.0003])
    book.save(path)
    return path


def test_excel_import_preserves_readings_precision_and_formulas(tmp_path):
    result = import_excel(source(tmp_path/'input.xlsx'),tmp_path/'output.xlsx', strict=True)
    assert result['measurement_columns'] == 2
    book = load_workbook(tmp_path/'output.xlsx',data_only=True)
    assert book.sheetnames == ['Summary Report','Differences']
    assert book.active['G7'].value == 1.0001
    assert book.active['H7'].value == 1.0003
    assert book.active['I7'].value == pytest.approx(1.0002)
    assert book.active['G7'].number_format == '0.0000'
    book.close()
    book = load_workbook(tmp_path/'output.xlsx')
    assert len(book.active._images) == 1
    assert len(book['Differences']._images) == 1
    assert book.active['D1'].value.startswith('Inspection Summary')
    assert book.active['I7'].data_type == 'f'
    book.close()


def test_excel_import_rejects_invalid_reading(tmp_path):
    with pytest.raises(ValueError,match='nonnumeric'):
        import_excel(source(tmp_path/'input.xlsx','broken'),tmp_path/'output.xlsx')
    assert not (tmp_path/'output.xlsx').exists()


def test_excel_import_unknown_layout_is_not_guessed(tmp_path):
    book = Workbook()
    book.active.append(['Invoice','Amount'])
    book.active.append(['INV1',100])
    book.save(tmp_path/'input.xlsx')
    with pytest.raises(ValueError,match='No recognizable measurement table'):
        import_excel(tmp_path/'input.xlsx',tmp_path/'output.xlsx')


def test_repeated_headers_reuse_reading_columns(tmp_path):
    path = source(tmp_path/'input.xlsx')
    book = load_workbook(path)
    book.active.append([cell.value for cell in book.active[2]])
    book.active.append(['mm','Second shaft','Diameter',1,'+/-0.1',1.02,1.03])
    book.save(path)
    book.close()
    result = import_excel(path,tmp_path/'output.xlsx')
    assert result['measurement_columns'] == 2
    output = load_workbook(tmp_path/'output.xlsx',data_only=True)
    assert output.active['G8'].value == 1.02
    assert output.active['H8'].value == 1.03
    output.close()


def test_missing_formula_cache_is_not_imported_as_blank(tmp_path):
    path = source(tmp_path/'input.xlsx')
    book = load_workbook(path)
    book.active['F3'] = '=1+0.01'
    book.save(path)
    book.close()
    with pytest.raises(ValueError,match='no saved result'):
        import_excel(path,tmp_path/'output.xlsx')


def test_conflicting_explicit_limits_require_review(tmp_path):
    path = source(tmp_path/'input.xlsx')
    book = load_workbook(path)
    book.active['H2'],book.active['I2'] = 'Lower Limit','Upper Limit'
    book.active['H3'],book.active['I3'] = .5,1.5
    book.save(path)
    book.close()
    result = import_excel(path,tmp_path/'output.xlsx')
    assert any('conflict' in note for note in result['warnings'])
